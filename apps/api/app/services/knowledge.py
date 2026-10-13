import io
import re
import zipfile
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader
from sqlalchemy import delete, select, text

from app.agents.providers import provider
from app.auth.security import Actor
from app.core.config import settings
from app.core.errors import DomainError
from app.db import repository as repo
from app.db.session import SessionLocal
from app.models import entities as m
from app.schemas.contracts import Search


def extract(data: bytes, suffix: str) -> list[tuple[int | None, str]]:
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 300:
            raise ValueError("Encrypted or oversized PDF")
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                raise ValueError("Expanded DOCX exceeds limits")
        document = WordDocument(io.BytesIO(data))
        content = "\n".join(p.text for p in document.paragraphs)
        content += "\n" + "\n".join(
            " | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows
        )
        return [(None, content)]
    return [(None, data.decode("utf-8-sig"))]


def chunks(pages):
    for page_number, content in pages:
        normalized = re.sub(r"[ \t]+", " ", content.replace("\x00", "")).strip()
        for start in range(0, len(normalized), 1000):
            piece = normalized[start : start + 1200]
            if piece.strip():
                yield page_number, piece


def ingest(document_id: str):
    # One ingestion per document, with deletion serialized by the same row lock.
    with SessionLocal() as db:
        row = db.scalar(select(m.Document).where(m.Document.id == document_id).with_for_update())
        if not row or row.status == "ready":
            return
        row.status = "processing"
        try:
            data = (Path(settings().upload_dir) / row.storage_key).read_bytes()
            pieces = list(chunks(extract(data, Path(row.name).suffix.lower())))
            if not pieces or len(pieces) > 1000:
                raise ValueError("No text or excessive chunks")
            db.execute(delete(m.DocumentChunk).where(m.DocumentChunk.document_id == row.id))
            embeddings = []
            for start in range(0, len(pieces), 32):
                embeddings.extend(provider().embed([part for _, part in pieces[start : start + 32]]))
            for ordinal, ((page_number, content), embedding) in enumerate(
                zip(pieces, embeddings, strict=True)
            ):
                db.add(
                    m.DocumentChunk(
                        organization_id=row.organization_id,
                        document_id=row.id,
                        ordinal=ordinal,
                        page=page_number,
                        content=content,
                        embedding=embedding,
                    )
                )
            row.status = "ready"
            row.chunk_count = len(pieces)
            row.error_code = None
            db.commit()
        except Exception:
            db.rollback()
            row = db.get(m.Document, document_id)
            if row:
                row.status = "failed"
                row.error_code = "extraction_or_embedding_failed"
                db.commit()
            raise DomainError("ingestion_failed", "Document extraction or embedding failed.", 502) from None


def search(db, actor: Actor, body: Search):
    vector = provider().embed([body.query])[0]
    distance = m.DocumentChunk.embedding.cosine_distance(vector)
    query = (
        select(m.DocumentChunk, m.Document, distance.label("distance"))
        .join(m.Document, m.Document.id == m.DocumentChunk.document_id)
        .where(
            m.DocumentChunk.organization_id == actor.organization_id,
            m.Document.organization_id == actor.organization_id,
            m.Document.status == "ready",
            m.Document.embedding_provider == settings().ai_provider,
            m.Document.embedding_model
            == (settings().embedding_model if settings().ai_provider == "openai" else "lexical-hash-v1"),
        )
    )
    if body.document_id:
        repo.get(db, m.Document, actor, body.document_id)
        query = query.where(m.Document.id == body.document_id)
    # Exact tenant-filtered scan guarantees top-K within a tenant, avoiding ANN post-filter starvation.
    db.execute(text("SET LOCAL enable_indexscan = off"))
    results = db.execute(
        query.where(distance <= 1 - body.threshold).order_by(distance).limit(body.top_k)
    ).all()
    return {
        "hits": [
            {
                "chunk_id": chunk.id,
                "document_id": document.id,
                "document_name": document.name,
                "page": chunk.page,
                "content": chunk.content,
                "score": round(1 - float(dist), 4),
            }
            for chunk, document, dist in results
        ],
        "debug": {"provider": settings().ai_provider, "top_k": body.top_k, "threshold": body.threshold},
    }


def validate_upload(name: str, data: bytes):
    if Path(name).suffix.lower() not in {".pdf", ".txt", ".md", ".docx"}:
        raise DomainError("unsupported_format", "Upload PDF, TXT, Markdown or DOCX.", 415)
    if not data or len(data) > settings().max_upload_bytes:
        raise DomainError("upload_size", "Upload a nonempty file no larger than 10 MB.", 413)
