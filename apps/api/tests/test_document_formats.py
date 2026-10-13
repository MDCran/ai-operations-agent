import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.knowledge import extract


def test_docx_extracts_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Support policy: contact the operations team.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Hours"
    table.cell(0, 1).text = "09:00–17:00"
    stream = io.BytesIO()
    document.save(stream)
    pages = extract(stream.getvalue(), ".docx")
    assert pages[0][0] is None
    assert "operations team" in pages[0][1] and "09:00" in pages[0][1]


def test_pdf_preserves_page_numbers_and_rejects_encryption():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    stream = io.BytesIO()
    writer.write(stream)
    assert extract(stream.getvalue(), ".pdf") == [(1, "")]
    writer.encrypt("do-not-decrypt")
    encrypted = io.BytesIO()
    writer.write(encrypted)
    with pytest.raises(ValueError, match="Encrypted"):
        extract(encrypted.getvalue(), ".pdf")


def test_docx_expansion_limit():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * (31 * 1024 * 1024))
    with pytest.raises(ValueError, match="Expanded DOCX"):
        extract(stream.getvalue(), ".docx")
